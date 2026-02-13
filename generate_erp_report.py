#!/usr/bin/env python3
"""生成 ERP 分析方法+结果+图注 的科研论文格式 Word 文档"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

doc = Document()

# ==================== 全局样式设置 ====================
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(6)

# 标题样式
for level in [1, 2, 3]:
    h_style = doc.styles[f'Heading {level}']
    h_font = h_style.font
    h_font.name = 'Times New Roman'
    h_font.bold = True
    h_font.color.rgb = RGBColor(0, 0, 0)
    h_style.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    if level == 1:
        h_font.size = Pt(16)
    elif level == 2:
        h_font.size = Pt(14)
    else:
        h_font.size = Pt(12)


def add_para(text, bold=False, italic=False, size=None, align=None, indent_first=True):
    """添加段落"""
    p = doc.add_paragraph()
    if indent_first:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    return p


def add_figure_note(fig_num, title, description):
    """添加图注（APA格式）"""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(12)
    # "图X" 加粗斜体
    run1 = p.add_run(f'图{fig_num}')
    run1.bold = True
    run1.italic = True
    run1.font.name = 'Times New Roman'
    run1.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run1.font.size = Pt(12)
    # 空格
    p.add_run(' ')
    # 标题（斜体）
    run2 = p.add_run(title)
    run2.italic = True
    run2.font.name = 'Times New Roman'
    run2.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run2.font.size = Pt(12)
    # 换行 + 描述（正体）
    desc_p = doc.add_paragraph()
    desc_p.paragraph_format.first_line_indent = Cm(0.74)
    desc_run = desc_p.add_run(description)
    desc_run.font.name = 'Times New Roman'
    desc_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    desc_run.font.size = Pt(10)
    return p


# ==================== 文档标题 ====================
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_p.add_run('闪烁光视觉搜索任务的ERP分析报告')
title_run.bold = True
title_run.font.size = Pt(22)
title_run.font.name = 'Times New Roman'
title_run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

subtitle_p = doc.add_paragraph()
subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run = subtitle_p.add_run('方法、结果与图注')
subtitle_run.font.size = Pt(14)
subtitle_run.font.name = 'Times New Roman'
subtitle_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
subtitle_run.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph()  # 空行

# ================================================================
#                          第一部分：方法
# ================================================================
doc.add_heading('1 方法', level=1)

# 1.1 被试
doc.add_heading('1.1 被试', level=2)
add_para(
    '本研究共招募10名被试参与实验。所有被试视力或矫正视力正常，无神经或精神疾病史。'
    '实验前所有被试签署知情同意书。'
)

# 1.2 实验设计
doc.add_heading('1.2 实验设计与刺激材料', level=2)
add_para(
    '本实验采用单因素三水平被试内设计，自变量为刺激类型（A刺激、B刺激、C刺激），'
    '分别以事件标记11、21、31进行编码。实验任务为闪烁光条件下的视觉搜索任务。'
    '每个试次（trial）的流程为：首先呈现刺激类型线索（标记11/21/31），'
    '随后呈现视觉搜索画面，其中有目标存在的试次以标记41编码，'
    '无目标存在的试次以标记42编码。被试需判断搜索画面中是否存在目标并做出按键反应，'
    '反应正确以标记12编码，反应错误以标记22编码，无反应以标记32编码。'
)

# 1.3 EEG记录
doc.add_heading('1.3 脑电记录', level=2)
add_para(
    '使用脑电采集系统连续记录头皮脑电信号（EEG），电极布置依照国际10-20系统扩展的电极位置，'
    '共59个有效记录通道。记录过程中以在线参考电极为参考，接地电极放置于前额。'
    '采集过程中阻抗保持在5 kΩ以下。'
)

# 1.4 EEG预处理
doc.add_heading('1.4 脑电预处理', level=2)
add_para(
    '离线分析使用MATLAB 2023及EEGLAB工具箱进行。预处理步骤包括：'
    '（1）带通滤波；（2）坏导插值；（3）重参考（如全脑平均参考）；'
    '（4）独立成分分析（ICA）去除眼电、肌电等伪迹成分；'
    '（5）以每个事件标记为锁时点进行分段（epoching）。'
    '预处理完成后，数据以EEGLAB的.set格式保存，每个被试的数据文件包含约950~1600个分段（epoch），'
    '涵盖所有事件类型。'
)

# 1.5 ERP分析
doc.add_heading('1.5 事件相关电位（ERP）分析', level=2)

doc.add_heading('1.5.1 条件提取与试次序列识别', level=3)
add_para(
    '由于预处理阶段以所有事件类型为锁时点进行了分段，本研究采用基于epoch时间顺序的条件识别策略：'
    '对于每个以标记41（有目标）或42（无目标）为锁时点的epoch，'
    '首先在该epoch的事件列表中查找潜伏期大于0的反应标记（12/22/32）以确定反应类型；'
    '其次，按epoch时间顺序向前回溯，查找前序epoch的锁时事件中最近的刺激类型标记（11/21/31），'
    '以确定当前试次的刺激条件。'
    '通过上述策略，每个有效epoch被分配一个条件编码。'
)
add_para(
    '最终提取的实验条件如下：'
)
add_para('有目标条件（事件41）：', bold=True, indent_first=True)
add_para('• 条件1：A刺激 + 有目标 + 正确反应（编码101）', indent_first=True)
add_para('• 条件2：B刺激 + 有目标 + 正确反应（编码201）', indent_first=True)
add_para('• 条件3：C刺激 + 有目标 + 正确反应（编码301）', indent_first=True)
add_para('无目标条件（事件42）：', bold=True, indent_first=True)
add_para('• 条件4：A刺激 + 无目标 + 正确反应（编码1101）', indent_first=True)
add_para('• 条件5：B刺激 + 无目标 + 正确反应（编码1201）', indent_first=True)
add_para('• 条件6：C刺激 + 无目标 + 正确反应（编码1301）', indent_first=True)

doc.add_heading('1.5.2 ERP平均与低通滤波', level=3)
add_para(
    '以搜索画面出现时刻（标记41或42）作为时间零点，分析窗口为刺激前100 ms至刺激后800 ms（−100~800 ms）。'
    '以刺激前100 ms（−100~0 ms）作为基线进行基线校正。'
    '对每个被试，分别对各条件下的有效epoch进行叠加平均，'
    '得到个体水平的ERP波形。'
)
add_para(
    '为去除平均波形中的高频噪声并使ERP成分更加清晰，'
    '对叠加平均后的波形施加4阶Butterworth低通滤波（截止频率30 Hz），'
    '采用零相位滤波方式（MATLAB filtfilt函数），以避免滤波引入的相位延迟对ERP成分潜伏期的影响。'
)

doc.add_heading('1.5.3 SSVEP减法：目标诱发ERP的提取', level=3)
add_para(
    '由于本实验采用闪烁光范式，被试在有目标和无目标条件下均受到稳态视觉诱发电位（Steady-State '
    'Visually Evoked Potential, SSVEP）的影响。为消除SSVEP干扰并提取纯粹由目标图像诱发的ERP成分，'
    '采用减法策略：对每个被试、每种刺激条件、每个电极和每个时间点，'
    '将无目标条件的平均ERP波形从有目标条件的平均ERP波形中减去：'
)
add_para(
    '目标诱发ERP = ERP（有目标, 正确） − ERP（无目标, 正确）',
    italic=True, indent_first=False
)
p_eq = doc.paragraphs[-1]
p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_para(
    '该方法基于以下假设：SSVEP成分在有目标和无目标条件中是相同的（因为闪烁光刺激参数一致），'
    '而目标图像诱发的ERP成分（如N2pc、P3等）仅在有目标条件中出现。'
    '因此，差值波形即反映了目标加工的纯粹神经活动。'
)

doc.add_heading('1.5.4 感兴趣电极与ERP成分', level=3)
add_para(
    '选取以下6个感兴趣电极进行分析：Fz（额中线）、Cz（中央中线）、Pz（顶中线）、'
    'Oz（枕中线）、PO7（左侧顶枕区）、PO8（右侧顶枕区）。'
    '重点分析的ERP成分及其时间窗口如下：'
)
add_para('• P1成分：80~130 ms，峰值约100 ms，反映早期视觉加工', indent_first=True)
add_para('• N1成分：130~200 ms，峰值约170 ms，反映注意选择的早期阶段', indent_first=True)
add_para('• N2成分：150~350 ms，峰值约250 ms，反映认知控制和注意分配', indent_first=True)
add_para('• P3成分：250~600 ms，峰值约400 ms，反映目标检测和工作记忆更新', indent_first=True)

doc.add_heading('1.5.5 峰值与平均振幅测量', level=3)
add_para(
    'N2成分的峰值振幅定义为150~350 ms时间窗内的最小值（负波峰），'
    'P3成分的峰值振幅定义为250~600 ms时间窗内的最大值（正波峰），'
    '相应的峰值潜伏期为该极值对应的时间点。'
    '此外，为提高测量的稳健性，还计算了N2（250±50 ms，即200~300 ms）'
    '和P3（400±100 ms，即300~500 ms）时间窗口内的平均振幅，用于后续统计分析。'
)

doc.add_heading('1.5.6 统计分析', level=3)
add_para(
    '对三种刺激条件（A、B、C）间的ERP差异进行单因素重复测量方差分析（repeated-measures ANOVA）。'
    '此外，为探索条件效应的时间动态特征，对各感兴趣电极的每个时间点分别进行'
    '逐时间点重复测量方差分析（point-by-point repeated-measures ANOVA），'
    '并采用错误发现率（False Discovery Rate, FDR）方法对多重比较进行校正（α = .05）。'
)

# ================================================================
#                          第二部分：结果
# ================================================================
doc.add_heading('2 结果', level=1)

doc.add_heading('2.1 行为数据概况', level=2)
add_para(
    '10名被试均完成了实验任务。在有目标条件下，各被试的正确反应试次数分布如下：'
    'A刺激条件平均约43个epoch（范围：36~51），'
    'B刺激条件平均约38个epoch（范围：25~67），'
    'C刺激条件平均约38个epoch（范围：22~62）。'
    '少量试次因预处理中的伪迹剔除或条件识别失败而被排除。'
)

doc.add_heading('2.2 原始ERP波形（有目标条件）', level=2)
add_para(
    '图1展示了三种刺激条件（A、B、C）在6个感兴趣电极上的组平均ERP波形。'
    '在所有电极上均可观察到典型的视觉ERP成分序列：'
    '刺激后约100 ms出现P1正波，约170 ms出现N1负波，'
    '约200~300 ms出现N2负波，约300~500 ms出现P3正波。'
    '在枕区和顶枕区电极（Oz、PO7、PO8）上，早期成分（P1、N1）较为突出；'
    '在中线电极（Cz、Pz）上，晚期成分（N2、P3）的条件间差异更为明显。'
)

add_figure_note(1, '三种刺激条件组平均ERP波形（正确反应）',
    '各子图分别对应一个感兴趣电极（Fz、Cz、Pz、Oz、PO7、PO8），'
    '蓝线代表A刺激条件，绿线代表B刺激条件，红线代表C刺激条件。'
    '波形为10名被试的组平均结果，经30 Hz低通滤波平滑处理。'
    'x轴为相对于搜索画面出现（标记41）的时间（ms），y轴为振幅（μV），负极朝上。')

doc.add_heading('2.3 正确与错误反应的ERP对比', level=2)
add_para(
    '图2展示了正确反应与错误反应条件下的ERP波形对比。'
    '在多个电极上，正确反应条件（实线）与错误反应条件（虚线）在N2和P3成分的振幅上存在差异，'
    '提示认知加工过程中正确与错误试次间的神经活动差异。'
    '但由于错误反应的试次数较少（平均每条件约5~15个epoch），'
    '错误条件的波形信噪比较低，后续统计分析主要聚焦于正确反应条件。'
)

add_figure_note(2, '正确反应与错误反应的ERP波形对比',
    '各子图分别对应一个感兴趣电极。实线代表正确反应条件，虚线代表错误反应条件，'
    '颜色编码与图1一致（蓝=A，绿=B，红=C）。'
    '波形为10名被试的组平均结果，经30 Hz低通滤波。y轴负极朝上。')

doc.add_heading('2.4 条件间差异波', level=2)
add_para(
    '图3展示了条件间差异波（B−A和C−A），用于直观反映刺激条件间的ERP差异。'
    '差异波在特定时间窗口偏离零线，提示相应时段内不同刺激条件诱发了不同强度的神经加工。'
)

add_figure_note(3, '条件间差异波（B−A和C−A）',
    '各子图分别对应一个感兴趣电极。绿线代表B条件减去A条件的差异波，'
    '红线代表C条件减去A条件的差异波。差异波偏离零线表明两种条件间ERP存在差异。'
    '波形为组平均结果，经30 Hz低通滤波。y轴负极朝上。')

doc.add_heading('2.5 ERP成分地形图', level=2)
add_para(
    '图4~6展示了关键ERP成分在不同时间窗口的头皮电压分布地形图。'
    'P1成分（100 ms）主要分布于枕区，N1成分（170 ms）分布于枕区和顶枕区，'
    '符合视觉早期加工的典型分布特征。N2成分（250 ms）的负波活动分布较广，'
    '在额中区和中央区较为显著。P3成分（400 ms）的正波活动主要集中在顶区和中央顶区，'
    '与目标检测和注意分配的经典分布一致。'
)
add_para(
    '图5和图6分别展示了三种刺激条件在N2和P3时间窗口内的平均地形图，'
    '可观察到不同刺激条件在这些成分上的空间分布差异。'
)

add_figure_note(4, '各ERP成分总平均地形图',
    '从左至右依次为P1（100 ms）、N1（170 ms）、N2（250 ms）、P3（400 ms）时间点的头皮电压分布。'
    '地形图为三种正确反应条件和10名被试的总平均。颜色表示电压值（μV），暖色代表正值，冷色代表负值。')

add_figure_note(5, '各条件N2时间窗口平均地形图',
    '从左至右分别为A、B、C三种刺激条件在N2时间窗口（250±50 ms，即200~300 ms）内的平均头皮电压分布。'
    '数据为10名被试的组平均。')

add_figure_note(6, '各条件P3时间窗口平均地形图',
    '从左至右分别为A、B、C三种刺激条件在P3时间窗口（400±100 ms，即300~500 ms）内的平均头皮电压分布。'
    '数据为10名被试的组平均。')

doc.add_heading('2.6 N2和P3振幅的条件间比较', level=2)
add_para(
    '图7和图8分别展示了各电极上N2和P3平均振幅的柱状图。'
    '对于N2成分（200~300 ms平均振幅），在Cz电极上，A条件振幅约为−2.50 μV，'
    'B条件约为−3.31 μV，C条件约为−2.70 μV。'
    '对于P3成分（300~500 ms平均振幅），A条件在多数电极上的P3振幅略大于B和C条件。'
    '各电极的详细数值见导出的CSV文件。'
)

add_figure_note(7, '各电极N2平均振幅柱状图',
    '每个子图对应一个感兴趣电极，三个柱分别代表A（蓝）、B（绿）、C（红）刺激条件'
    '在N2时间窗口（250±50 ms）内的平均振幅。误差线表示标准误（SE）。'
    '数据为10名被试的组平均。')

add_figure_note(8, '各电极P3平均振幅柱状图',
    '每个子图对应一个感兴趣电极，三个柱分别代表A（蓝）、B（绿）、C（红）刺激条件'
    '在P3时间窗口（400±100 ms）内的平均振幅。误差线表示标准误（SE）。'
    '数据为10名被试的组平均。')

doc.add_heading('2.7 逐时间点方差分析', level=2)
add_para(
    '图9展示了各电极上逐时间点单因素重复测量方差分析的结果。'
    '上排为三种条件的组平均波形，下排为对应的p值曲线。'
    '灰色虚线标示未校正的显著性阈值（p = .05），红色实线标示FDR校正后的显著性阈值。'
    'p值曲线低于FDR阈值的时间区段内，三种刺激条件间的ERP振幅存在统计学显著差异。'
)

add_figure_note(9, '逐时间点重复测量方差分析结果',
    '上排各子图为三种刺激条件（A=蓝，B=绿，C=红）在各电极上的组平均ERP波形。'
    '下排各子图为对应电极上逐时间点单因素重复测量方差分析的p值曲线。'
    '灰色虚线 = p = .05（未校正），红色实线 = FDR校正阈值。'
    'p值低于红线的时间段表示三种条件间差异达到FDR校正后的显著性水平。'
    '分析基于经30 Hz低通滤波后的数据。')

doc.add_heading('2.8 SSVEP减法：目标诱发ERP', level=2)
add_para(
    '为消除闪烁光引起的稳态视觉诱发电位（SSVEP）干扰，将无目标条件的ERP从有目标条件的ERP中减去，'
    '得到目标诱发的纯ERP波形。'
)
add_para(
    '图10展示了SSVEP消除后的目标诱发ERP波形。与图1的原始波形相比，'
    '差值波形中SSVEP相关的周期性振荡被显著抑制，目标特异性的ERP成分更加清晰。'
    '在顶枕区电极（PO7、PO8）上可更清晰地观察到N2pc等目标选择相关成分，'
    '在中线电极（Cz、Pz）上P3成分更加突出。'
)

add_figure_note(10, '目标诱发ERP波形（SSVEP已消除）',
    '各子图分别对应一个感兴趣电极。波形为有目标条件正确试次的ERP减去无目标条件正确试次的ERP后的差值波形。'
    '蓝线=A条件，绿线=B条件，红线=C条件。'
    '该减法消除了有目标和无目标条件中共有的SSVEP成分，保留了目标特异性的ERP活动。'
    '波形为10名被试的组平均结果，经30 Hz低通滤波。y轴负极朝上。')

add_para(
    '图11将有目标、无目标和差值（目标诱发ERP）三条波形叠加在同一图中，'
    '直观展示了SSVEP减法的效果。可以看到有目标和无目标条件的波形在早期（0~150 ms）较为接近'
    '（因为SSVEP成分在两个条件中基本一致），而在150 ms之后出现分离，'
    '差值波形（黑线）反映了目标加工引起的额外神经活动。'
)

add_figure_note(11, '有目标、无目标与目标诱发ERP的对比',
    '各子图分别对应一个感兴趣电极（所有刺激条件合并平均）。'
    '红线=有目标条件ERP，蓝线=无目标条件ERP，黑色粗线=差值（目标诱发ERP = 有目标 − 无目标）。'
    '波形为10名被试的组平均结果。y轴负极朝上。'
    '两条件在早期阶段波形接近，说明SSVEP成分在有/无目标条件中基本一致，减法策略有效。')

add_para(
    '图12展示了目标诱发ERP在P1、N1、N2、P3各时间点的头皮地形图，'
    '反映了消除SSVEP后目标加工相关神经活动的空间分布特征。'
)

add_figure_note(12, '目标诱发ERP地形图（SSVEP消除后）',
    '从左至右依次为P1（100 ms）、N1（170 ms）、N2（250 ms）、P3（400 ms）时间点的'
    '目标诱发ERP（有目标−无目标）头皮电压分布。'
    '数据为三种条件和10名被试的总平均。颜色表示差值电压（μV）。')

# ================================================================
#                          附录：分析参数汇总
# ================================================================
doc.add_heading('附录：分析参数汇总', level=1)

# 参数表格
table = doc.add_table(rows=13, cols=2, style='Table Grid')
table.columns[0].width = Inches(2.5)
table.columns[1].width = Inches(4.0)

params = [
    ('参数', '设置值'),
    ('被试数', '10'),
    ('分析时间窗口', '−100 ~ 800 ms'),
    ('基线校正', '−100 ~ 0 ms'),
    ('低通滤波截止频率', '30 Hz（4阶Butterworth，零相位）'),
    ('感兴趣电极', 'Fz, Cz, Pz, Oz, PO7, PO8'),
    ('N2搜索窗口', '150 ~ 350 ms'),
    ('P3搜索窗口', '250 ~ 600 ms'),
    ('N2平均振幅窗口', '200 ~ 300 ms（250 ± 50 ms）'),
    ('P3平均振幅窗口', '300 ~ 500 ms（400 ± 100 ms）'),
    ('统计方法', '单因素重复测量ANOVA + FDR校正'),
    ('SSVEP消除', '有目标ERP − 无目标ERP'),
    ('分析工具', 'MATLAB 2023 + EEGLAB 2023.1'),
]

for row_idx, (col1, col2) in enumerate(params):
    cell0 = table.rows[row_idx].cells[0]
    cell1 = table.rows[row_idx].cells[1]
    cell0.text = col1
    cell1.text = col2
    # 表头加粗
    if row_idx == 0:
        for cell in [cell0, cell1]:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

# 设置表格字体
for row in table.rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.first_line_indent = Cm(0)
            for run in paragraph.runs:
                run.font.size = Pt(10)
                run.font.name = 'Times New Roman'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# ==================== 保存 ====================
output_path = '/workspace/ERP分析报告_方法结果图注.docx'
doc.save(output_path)
print(f'文档已保存至: {output_path}')
